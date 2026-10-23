#!/usr/bin/env python3
"""Build the revised Upper Sixth ICT book from one shared, reviewable source.

No table-based layout. DOCX and PDF consume the same parsed blocks, objectives,
code and original figures. PDF pagination is independent of Word's pagination.
Run from any directory: .venv/bin/python scripts/build_ict_notes.py
"""
from __future__ import annotations

import argparse
from dataclasses import dataclass
from html import escape
from datetime import datetime, timezone
import json
import os
from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from PIL import Image as PILImage
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    BaseDocTemplate, Frame, HRFlowable, Image, KeepTogether, PageBreak,
    PageTemplate, Paragraph, Preformatted, Spacer,
)
from reportlab.platypus.tableofcontents import TableOfContents

from ict_diagrams import draw_all

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "notes/upper-sixth-ict"
BASENAME = "Upper Sixth - ICT - First Term Teaching Notes"
INK = "253448"
NAVY = "192D46"
BLUE = "2459A6"
GREY = "526276"
PALE = "EEF4FC"
REVISED = "23 October 2026"
FONT_DIR = Path(os.environ.get("ICT_NOTES_FONT_DIR", "/usr/share/fonts/truetype/dejavu"))
TOKEN = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\(https?://[^)]+\)|\*[^*\n]+\*)")
LESSON = re.compile(r"^Lesson (\d+) — ")
IMAGE = re.compile(r"^!\[([^\]]+)\]\(([^)]+)\)$")


@dataclass
class Block:
    kind: str
    text: str = ""
    level: int = 0
    target: str = ""
    marker: str = ""
    anchor: str = ""
    toc_level: int = -1
    new_page: bool = False


def parse_markdown(text: str) -> list[Block]:
    lines = text.splitlines()
    result = []
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line.startswith("```"):
            info = line[3:].strip().split(maxsplit=1)
            language = info[0] if info else ""
            filename = info[1] if len(info) == 2 else ""
            code = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            if i == len(lines):
                raise ValueError("Unclosed code fence")
            result.append(Block("code", "\n".join(code), target=filename, marker=language))
            i += 1
            continue
        image = IMAGE.match(line)
        if image:
            result.append(Block("image", image[1], target=image[2]))
            i += 1
            continue
        heading = re.match(r"^(#{1,3}) (.+)$", line)
        if heading:
            result.append(Block("heading", heading[2], level=len(heading[1])))
            i += 1
            continue
        item = re.match(r"^(\d+\.|-) (.+)$", line)
        if item:
            result.append(Block("item", item[2], marker=item[1]))
            i += 1
            continue
        if line.startswith("|"):
            raise ValueError("Table markup is intentionally not supported in the revised book")
        paragraph = [line]
        i += 1
        while i < len(lines) and lines[i].strip():
            s = lines[i].strip()
            if s.startswith(("# ", "## ", "### ", "```", "![", "- ")) or re.match(r"^\d+\. ", s):
                break
            paragraph.append(s)
            i += 1
        result.append(Block("paragraph", " ".join(paragraph)))
    return result


def load_book() -> tuple[dict, list[Block]]:
    curriculum = json.loads((SOURCE / "curriculum.json").read_text(encoding="utf-8"))
    by_number = {entry["number"]: entry for entry in curriculum["lessons"]}
    paths = [SOURCE / "front-matter.md"]
    paths += [SOURCE / part["file"] for part in curriculum["parts"]]
    paths += [SOURCE / "revision-and-sources.md"]
    result = []
    seen = []
    part_follows = False
    for path in paths:
        for block in parse_markdown(path.read_text(encoding="utf-8")):
            if block.kind == "heading":
                if block.level == 1:
                    block.toc_level = 0
                    block.new_page = True
                    part = re.match(r"Part (\d+) —", block.text)
                    block.anchor = f"part-{part[1]}" if part else f"section-{len(result)}"
                    part_follows = bool(part)
                elif block.level == 2:
                    match = LESSON.match(block.text)
                    evaluation = re.match(r"Evaluation (\d+) —", block.text)
                    if match or evaluation:
                        block.toc_level = 1
                        block.new_page = not part_follows
                        block.anchor = (f"lesson-{match[1]}" if match
                                        else f"evaluation-{evaluation[1]}")
                        part_follows = False
                    if match:
                        number = int(match[1])
                        seen.append(number)
                        entry = by_number[number]
                        result.append(block)
                        detail = (f"Week {entry['week']} · Supplied progression: PDF page "
                                  f"{entry['progression_page']}")
                        if entry["type"] == "integration":
                            detail += " · Locally authored integration task"
                        result.append(Block("meta", detail))
                        if entry["objectives"]:
                            result.append(Block("outcome_title", "Learning outcomes — supplied progression"))
                            result.extend(Block("objective", obj, marker="•") for obj in entry["objectives"])
                        continue
            result.append(block)
    if seen != list(range(1, 47)):
        raise ValueError(f"Expected Lessons 1–46 in order, found {seen}")
    # Keep short exercise sets and integration-review tails together. A final
    # answer or review paragraph should not create a mostly empty printed page.
    # These markers affect layout only; no teaching text is added or removed.
    grouped = []
    group_type = ""
    for block in result:
        if block.kind == "heading":
            starts_exercises = block.text == "Check your understanding"
            starts_review = block.text.startswith("Assessment criteria —")
            closes_group = ((group_type == "exercises" and block.text != "Answer guidance")
                            or (group_type == "review" and block.level <= 2))
            if group_type and closes_group:
                grouped.append(Block("group_end"))
                group_type = ""
            if starts_exercises or starts_review:
                grouped.append(Block("group_start"))
                group_type = "exercises" if starts_exercises else "review"
        grouped.append(block)
    if group_type:
        grouped.append(Block("group_end"))
    return curriculum, grouped


def inline_tokens(text):
    for token in TOKEN.split(text):
        if not token:
            continue
        if token.startswith("**") and token.endswith("**"):
            yield "bold", token[2:-2], ""
        elif token.startswith("`") and token.endswith("`"):
            yield "code", token[1:-1], ""
        elif token.startswith("[") and "](" in token:
            label, url = token[1:].split("](", 1)
            yield "link", label, url[:-1]
        elif token.startswith("*") and token.endswith("*"):
            yield "italic", token[1:-1], ""
        else:
            yield "text", token, ""


def pdf_inline(text):
    parts = []
    for kind, value, url in inline_tokens(text):
        value = escape(value)
        if kind == "bold":
            value = f"<b>{value}</b>"
        elif kind == "italic":
            value = f"<i>{value}</i>"
        elif kind == "code":
            value = f'<font name="BookMono" size="9.5">{value}</font>'
        elif kind == "link":
            value = f'<link href="{escape(url, quote=True)}" color="#{BLUE}">[{value}]</link>'
        parts.append(value)
    return "".join(parts)


def docx_link(paragraph, label, target, internal=False):
    link = OxmlElement("w:hyperlink")
    if internal:
        link.set(qn("w:anchor"), target)
    else:
        rel = paragraph.part.relate_to(target, RT.HYPERLINK, is_external=True)
        link.set(qn("r:id"), rel)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    props.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(underline)
    run.append(props)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    link.append(run)
    paragraph._p.append(link)


def word_inline(paragraph, text):
    for kind, value, url in inline_tokens(text):
        if kind == "link":
            docx_link(paragraph, f"[{value}]", url)
            continue
        run = paragraph.add_run(value)
        if kind == "bold":
            run.bold = True
        elif kind == "italic":
            run.italic = True
        elif kind == "code":
            run.font.name = "DejaVu Sans Mono"
            run.font.size = Pt(9.5)


def bookmark(paragraph, name, index):
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(index))
    start.set(qn("w:name"), name.replace("-", "_"))
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(index))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def shade_paragraph(paragraph, fill=PALE):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    paragraph._p.get_or_add_pPr().append(shd)


def make_docx(curriculum, blocks, destination):
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.left_margin = section.right_margin = Cm(1.9)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(1.9)
    section.header_distance = section.footer_distance = Cm(0.85)
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "DejaVu Serif"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.line_spacing = 1.24
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.widow_control = True
    for level, size in [(1, 23), (2, 18), (3, 12.3)]:
        s = styles[f"Heading {level}"]
        s.font.name = "DejaVu Sans"
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.color.rgb = RGBColor.from_string(BLUE if level == 3 else NAVY)
        s.paragraph_format.space_before = Pt(13 if level > 1 else 4)
        s.paragraph_format.space_after = Pt(6)
        s.paragraph_format.keep_with_next = True
    for name in ["Book Meta", "Book Objective", "Book Code", "Book Caption", "Book Contents"]:
        if name not in styles:
            styles.add_style(name, 1)
    styles["Book Meta"].font.name = "DejaVu Sans"
    styles["Book Meta"].font.size = Pt(9)
    styles["Book Meta"].font.color.rgb = RGBColor.from_string(GREY)
    styles["Book Meta"].paragraph_format.space_after = Pt(7)
    styles["Book Objective"].font.size = Pt(10)
    styles["Book Objective"].paragraph_format.left_indent = Cm(0.35)
    styles["Book Objective"].paragraph_format.first_line_indent = Cm(-0.25)
    styles["Book Objective"].paragraph_format.space_after = Pt(3)
    styles["Book Objective"].paragraph_format.line_spacing = 1.15
    styles["Book Code"].font.name = "DejaVu Sans Mono"
    styles["Book Code"].font.size = Pt(8.7)
    styles["Book Code"].paragraph_format.line_spacing = 1.0
    styles["Book Code"].paragraph_format.space_after = Pt(0)
    styles["Book Code"].paragraph_format.space_before = Pt(0)
    styles["Book Code"].paragraph_format.left_indent = Cm(0.2)
    styles["Book Code"].paragraph_format.keep_with_next = False
    styles["Book Caption"].font.name = "DejaVu Sans"
    styles["Book Caption"].font.size = Pt(9)
    styles["Book Caption"].font.color.rgb = RGBColor.from_string(GREY)
    styles["Book Caption"].paragraph_format.space_after = Pt(10)
    styles["Book Caption"].paragraph_format.line_spacing = 1.1
    styles["Book Contents"].font.name = "DejaVu Sans"
    styles["Book Contents"].font.size = Pt(10)
    styles["Book Contents"].paragraph_format.space_after = Pt(4)
    styles["Book Contents"].paragraph_format.line_spacing = 1.12
    header = section.header.paragraphs[0]
    header.text = "UPPER SIXTH  ·  ICT                                      FIRST TERM  ·  2026–2027"
    header.style = styles["Book Meta"]
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.style = styles["Book Meta"]
    footer.add_run("Learning notes  ·  ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    document.core_properties.title = curriculum["title"]
    document.core_properties.author = "Upper Sixth ICT teaching resource"
    document.core_properties.created = datetime(2026, 10, 23, tzinfo=timezone.utc)
    document.core_properties.modified = datetime(2026, 10, 23, tzinfo=timezone.utc)
    document.core_properties.subject = "First-term ICT explanations and practical learning, Lessons 1–46"
    document.core_properties.keywords = "Upper Sixth, ICT, Cameroon, learning notes, 2026–2027"
    document.core_properties.comments = "Based on the supplied progression; not an official GCE publication."

    p = document.add_paragraph()
    p.paragraph_format.space_before = Pt(75)
    run = p.add_run("INFORMATION & COMMUNICATION TECHNOLOGY")
    run.font.name = "DejaVu Sans"
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor.from_string(BLUE)
    p = document.add_paragraph()
    r = p.add_run("UPPER SIXTH")
    r.font.name = "DejaVu Sans"
    r.font.size = Pt(22)
    r.font.color.rgb = RGBColor.from_string(NAVY)
    p = document.add_paragraph()
    r = p.add_run("ICT")
    r.font.name = "DejaVu Sans"
    r.font.size = Pt(62)
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(BLUE)
    p = document.add_paragraph()
    r = p.add_run("First Term Learning Notes")
    r.font.name = "DejaVu Sans"
    r.font.size = Pt(25)
    r.font.color.rgb = RGBColor.from_string(NAVY)
    document.add_paragraph("Explanations · worked examples · labelled diagrams\nGuided practice · independent questions · answer guidance")
    p = document.add_paragraph("2026–2027  |  Weeks 1–13  |  Lessons 1–46")
    p.paragraph_format.space_before = Pt(25)
    document.add_paragraph("35 original figures · 9 integration activities · 2 classroom evaluations")
    document.add_paragraph("A teaching resource based on the supplied progression.\nNot an official GCE Board publication.", style="Book Meta")
    document.add_paragraph(f"Revised {REVISED}", style="Book Meta")
    document.add_page_break()
    document.add_heading("Contents", 1)
    document.add_paragraph("Click a lesson title or use Word's heading navigation. The PDF has its own paginated contents.", style="Book Meta")
    for block in blocks:
        if block.toc_level >= 0:
            p = document.add_paragraph(style="Book Contents")
            p.paragraph_format.left_indent = Cm(0.45 if block.toc_level else 0)
            if block.toc_level == 0:
                p.paragraph_format.space_before = Pt(6)
            docx_link(p, block.text, block.anchor.replace("-", "_"), internal=True)

    bookmark_index = 0
    in_exercises = False
    for block in blocks:
        if block.kind == "group_start":
            in_exercises = True
            continue
        if block.kind == "group_end":
            document.paragraphs[-1].paragraph_format.keep_with_next = False
            in_exercises = False
            continue
        paragraph_count = len(document.paragraphs)
        if block.kind == "heading":
            p = document.add_heading("", block.level)
            if block.new_page:
                p.paragraph_format.page_break_before = True
            word_inline(p, block.text)
            if block.anchor:
                bookmark_index += 1
                bookmark(p, block.anchor, bookmark_index)
        elif block.kind == "meta":
            document.add_paragraph(block.text, style="Book Meta")
        elif block.kind == "outcome_title":
            p = document.add_paragraph(style="Book Meta")
            p.add_run(block.text).bold = True
            p.paragraph_format.keep_with_next = True
        elif block.kind == "objective":
            p = document.add_paragraph(style="Book Objective")
            p.add_run("•  " + block.text)
        elif block.kind == "paragraph":
            word_inline(document.add_paragraph(), block.text)
        elif block.kind == "item":
            p = document.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.45)
            p.paragraph_format.first_line_indent = Cm(-0.45)
            p.add_run(("•" if block.marker == "-" else block.marker) + "  ")
            word_inline(p, block.text)
        elif block.kind == "code":
            if block.target:
                p = document.add_paragraph(f"Example file: {block.target}", style="Book Meta")
                p.paragraph_format.keep_with_next = True
            for line in block.text.splitlines():
                p = document.add_paragraph(style="Book Code")
                p.add_run(line or " ")
                shade_paragraph(p, "F4F6F8")
            document.add_paragraph().paragraph_format.space_after = Pt(2)
        elif block.kind == "image":
            path = SOURCE / block.target
            p = document.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_after = Pt(3)
            picture = p.add_run().add_picture(str(path), width=Cm(17.1))
            picture._inline.docPr.set("descr", block.text)
            picture._inline.docPr.set("title", block.text.split(" — ")[0])
            word_inline(document.add_paragraph(style="Book Caption"), block.text)
        if in_exercises:
            for paragraph in document.paragraphs[paragraph_count:]:
                paragraph.paragraph_format.keep_with_next = True
    document.save(destination)


class BookPDF(BaseDocTemplate):
    def afterFlowable(self, flowable):
        level = getattr(flowable, "book_toc_level", -1)
        if level >= 0:
            key, title = flowable.book_anchor, flowable.book_title
            self.canv.bookmarkPage(key)
            self.canv.addOutlineEntry(title, key, level=level, closed=False)
            self.notify("TOCEntry", (level, escape(title), self.page, key))


def register_pdf_fonts():
    for name, file in [("BookSerif", "DejaVuSerif.ttf"), ("BookSerif-Bold", "DejaVuSerif-Bold.ttf"),
                        ("BookSans", "DejaVuSans.ttf"), ("BookSans-Bold", "DejaVuSans-Bold.ttf"),
                        ("BookMono", "DejaVuSansMono.ttf")]:
        pdfmetrics.registerFont(TTFont(name, str(FONT_DIR / file)))
    for family in ["BookSerif", "BookSans"]:
        pdfmetrics.registerFontFamily(family, normal=family, bold=family + "-Bold",
                                     italic=family, boldItalic=family + "-Bold")
    pdfmetrics.registerFontFamily("BookMono", normal="BookMono", bold="BookMono",
                                 italic="BookMono", boldItalic="BookMono")


def make_pdf(curriculum, blocks, destination):
    register_pdf_fonts()
    margin = 18.5 * mm
    doc = BookPDF(str(destination), pagesize=A4, leftMargin=margin, rightMargin=margin,
                  topMargin=22 * mm, bottomMargin=19 * mm,
                  title=curriculum["title"], author="ICT classroom learning resource",
                  subject="Lessons 1–46, based on the supplied 2026–2027 progression",
                  pageCompression=1, allowSplitting=1)
    frame = Frame(doc.leftMargin, doc.bottomMargin, doc.width, doc.height,
                  leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0, id="book")

    def page(canvas, document):
        canvas.saveState()
        canvas.setFont("BookSans", 8)
        canvas.setFillColor(colors.HexColor("#" + GREY))
        if document.page > 1:
            canvas.drawString(margin, A4[1] - 12 * mm, "UPPER SIXTH  ·  ICT")
            canvas.drawRightString(A4[0] - margin, A4[1] - 12 * mm, "FIRST TERM  ·  2026–2027")
            canvas.setStrokeColor(colors.HexColor("#D6DFEB"))
            canvas.setLineWidth(0.4)
            canvas.line(margin, A4[1] - 15 * mm, A4[0] - margin, A4[1] - 15 * mm)
        canvas.drawString(margin, 10 * mm, "Learning notes  ·  Based on the supplied progression")
        canvas.drawRightString(A4[0] - margin, 10 * mm, str(document.page))
        canvas.restoreState()

    doc.addPageTemplates(PageTemplate(id="book", frames=[frame], onPage=page))
    c = lambda h: colors.HexColor("#" + h)
    body = ParagraphStyle("Body", fontName="BookSerif", fontSize=10.7, leading=15,
                          textColor=c(INK), spaceAfter=7, allowWidows=0, allowOrphans=0,
                          splitLongWords=1, uriWasteReduce=0.15)
    h1 = ParagraphStyle("Part", fontName="BookSans-Bold", fontSize=22, leading=28,
                        textColor=c(NAVY), spaceAfter=13, keepWithNext=True)
    h2 = ParagraphStyle("Lesson", fontName="BookSans-Bold", fontSize=17.5, leading=22.5,
                        textColor=c(NAVY), spaceBefore=8, spaceAfter=8, keepWithNext=True)
    h3 = ParagraphStyle("Topic", fontName="BookSans-Bold", fontSize=11.8, leading=16,
                        textColor=c(BLUE), spaceBefore=10, spaceAfter=5, keepWithNext=True)
    meta = ParagraphStyle("Meta", parent=body, fontName="BookSans", fontSize=8.8,
                          leading=12, textColor=c(GREY), spaceAfter=7)
    outcome_title = ParagraphStyle("OutcomeTitle", parent=meta, fontName="BookSans-Bold", keepWithNext=True)
    objective = ParagraphStyle("Objective", parent=body, fontSize=9.6, leading=13,
                               leftIndent=10, firstLineIndent=0, spaceAfter=3,
                               bulletFontName="BookSans", bulletFontSize=8.5)
    item = ParagraphStyle("Item", parent=body, leftIndent=14, firstLineIndent=0,
                          bulletIndent=0, spaceAfter=7)
    caption = ParagraphStyle("Caption", parent=meta, fontSize=8.5, leading=11.5, spaceAfter=10)
    code = ParagraphStyle("Code", fontName="BookMono", fontSize=8.3, leading=11.1,
                           leftIndent=8, rightIndent=8, backColor=c("F4F6F8"),
                           borderPadding=7, spaceBefore=3, spaceAfter=10)
    story = [Spacer(1, 40 * mm)]
    story.append(Paragraph("INFORMATION &amp; COMMUNICATION TECHNOLOGY", meta))
    story.append(Spacer(1, 8 * mm))
    story.append(Paragraph("UPPER SIXTH", ParagraphStyle("CoverClass", parent=h1, fontSize=21)))
    story.append(Paragraph("ICT", ParagraphStyle("CoverICT", parent=h1, fontSize=68, leading=78, textColor=c(BLUE))))
    story.append(Spacer(1, 7 * mm))
    story.append(Paragraph("First Term Learning Notes", ParagraphStyle("CoverTitle", parent=h1, fontSize=25, leading=32)))
    story.append(HRFlowable(width="100%", thickness=1.5, color=c(BLUE), spaceBefore=4, spaceAfter=14))
    story.append(Paragraph("Explanations · worked examples · labelled diagrams<br/>Guided practice · independent questions · answer guidance", body))
    story.append(Spacer(1, 14 * mm))
    story.append(Paragraph("<b>2026–2027</b>  |  Weeks 1–13  |  Lessons 1–46", body))
    story.append(Paragraph("35 original figures · 9 integration activities · 2 classroom evaluations", meta))
    story.append(Paragraph("A teaching resource based on the supplied progression.<br/>Not an official GCE Board publication.", meta))
    story.append(Paragraph(f"Revised {REVISED}", meta))
    story.append(PageBreak())
    story.append(Paragraph("Contents", h1))
    story.append(Paragraph("Page numbers and lesson titles are clickable in this PDF.", meta))
    toc = TableOfContents()
    toc.levelStyles = [
        ParagraphStyle("TOC0", fontName="BookSans-Bold", fontSize=10.1, leading=14,
                       textColor=c(NAVY), spaceBefore=7, leftIndent=0, firstLineIndent=0),
        ParagraphStyle("TOC1", fontName="BookSerif", fontSize=9.2, leading=12.5,
                       textColor=c(INK), spaceBefore=2, leftIndent=12, firstLineIndent=0),
    ]
    toc.dotsMinLevel = 0
    story.append(toc)
    exercise_group = None
    for block in blocks:
        if block.kind == "group_start":
            exercise_group = []
            continue
        if block.kind == "group_end":
            story.append(KeepTogether(exercise_group))
            exercise_group = None
            continue
        flowables = exercise_group if exercise_group is not None else story
        if block.kind == "heading":
            if block.new_page:
                flowables.append(PageBreak())
            p = Paragraph(pdf_inline(block.text), {1: h1, 2: h2, 3: h3}[block.level])
            if block.toc_level >= 0:
                p.book_toc_level = block.toc_level
                p.book_anchor = block.anchor
                p.book_title = block.text
            flowables.append(p)
        elif block.kind == "meta":
            flowables.append(Paragraph(pdf_inline(block.text), meta))
        elif block.kind == "outcome_title":
            flowables.append(Paragraph(pdf_inline(block.text), outcome_title))
        elif block.kind == "objective":
            flowables.append(Paragraph(pdf_inline(block.text), objective, bulletText="•"))
        elif block.kind == "paragraph":
            flowables.append(Paragraph(pdf_inline(block.text), body))
        elif block.kind == "item":
            flowables.append(Paragraph(pdf_inline(block.text), item,
                                   bulletText="•" if block.marker == "-" else block.marker))
        elif block.kind == "code":
            if block.target:
                label = Paragraph("Example file: " + escape(block.target), outcome_title)
                flowables.append(label)
            flowables.append(Preformatted(block.text, code, maxLineLength=88))
        elif block.kind == "image":
            path = SOURCE / block.target
            with PILImage.open(path) as im:
                width, height = im.size
            diagram = Image(str(path), width=doc.width, height=doc.width * height / width)
            diagram.hAlign = "CENTER"
            flowables.append(KeepTogether([Spacer(1, 5), diagram,
                                      Paragraph(pdf_inline(block.text), caption)]))
    doc.multiBuild(story, maxPasses=5)


def alignment_report(curriculum):
    lines = ["# Upper Sixth ICT — First-term alignment", "",
             "Primary source: `" + curriculum["primary_source"] + "`.", "",
             curriculum["source_note"], "", curriculum["objectives_note"], "",
             "Part numbers below are editorial reading sections, not official module numbers.", ""]
    for part in curriculum["parts"]:
        lines += [f"## Part {part['number']} — {part['title']}", ""]
        for entry in curriculum["lessons"]:
            if entry["part"] != part["number"]:
                continue
            lines += [f"### Lesson {entry['number']} — {entry['title']}", "",
                      f"Week {entry['week']} · supplied PDF page {entry['progression_page']} · "
                      f"reading source: [{part['file']}]({part['file']})", ""]
            if entry["objectives"]:
                lines += ["- " + objective for objective in entry["objectives"]] + [""]
            else:
                lines += ["Integration lesson: the sheet lists no separate objectives. The scenario, evidence requirements and classroom criteria are locally authored.", ""]
    lines += ["## Evaluation placement", "",
              "- Evaluation 1 and remediation follow Lesson 17 (Week 5).",
              "- Evaluation 2 and remediation follow Lesson 34 (Week 10), before parity and Internet history.",
              "- First term ends at Lesson 46. CSS from Lesson 47 onwards is not taught in this book.", ""]
    return "\n".join(lines)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir", type=Path, default=ROOT / "output")
    parser.add_argument("--skip-diagrams", action="store_true")
    args = parser.parse_args()
    if not args.skip_diagrams:
        draw_all()
    curriculum, blocks = load_book()
    args.output_dir.mkdir(parents=True, exist_ok=True)
    for block in blocks:
        if block.kind == "image" and not (SOURCE / block.target).is_file():
            raise FileNotFoundError(block.target)
        if block.kind == "code" and block.target:
            if Path(block.target).name != block.target or not block.target.endswith(".html"):
                raise ValueError(f"Example filename must be a local HTML basename: {block.target}")
            target = ROOT / "examples/upper-sixth-ict/site" / block.target
            target.parent.mkdir(parents=True, exist_ok=True)
            target.write_text(block.text + "\n", encoding="utf-8")
    (SOURCE / "alignment.md").write_text(alignment_report(curriculum), encoding="utf-8")
    word_path = args.output_dir / f"{BASENAME}.docx"
    pdf_path = args.output_dir / f"{BASENAME}.pdf"
    make_docx(curriculum, blocks, word_path)
    print(f"Wrote {word_path.relative_to(ROOT) if word_path.is_relative_to(ROOT) else word_path}", flush=True)
    make_pdf(curriculum, blocks, pdf_path)
    print(f"Wrote {pdf_path.relative_to(ROOT) if pdf_path.is_relative_to(ROOT) else pdf_path}", flush=True)
    print(f"{len(curriculum['lessons'])} lessons; "
          f"{sum(len(x['objectives']) for x in curriculum['lessons'])} progression outcomes; "
          f"{sum(b.kind == 'image' for b in blocks)} original figures; "
          f"{sum(len(b.text.split()) for b in blocks):,} exported words.")


if __name__ == "__main__":
    main()
